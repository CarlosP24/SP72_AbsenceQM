"""Build the ICTP talk deck on top of the SuperSemi template (June 2026).

The template's theme, master, layouts and logos are kept; its title and farewell
slides are edited in place and every other slide is rebuilt from the "Title Only"
layout. Text is native and editable (theme fonts set to Montserrat; Greek letters and
maths symbols, missing from Montserrat, go in Cambria Math). Figures and equations are
SVG pictures with a PNG fallback. Slides use Morph transitions (fade fallback) and
click-triggered fade-in builds, as in the earlier decks; speaker notes carry the
talking points and target times.

Run:  python slides/build_deck.py [output.pptx]
Needs python-pptx and lxml. Inputs: slides/figs, slides/equations, slides/assets
(template, photos and logos, not tracked), plots/figures (paper figures for backups).
"""
import copy
import itertools
import json
import pathlib
import re
import subprocess
import sys
import tempfile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
REPO = HERE.parent
FIGS, EQS, ASSETS = HERE / "figs", HERE / "equations", HERE / "assets"
TEMPLATE = ASSETS / "template_supersemi.pptx"
DEFAULT_OUT = (REPO.parents[1] / "Presentations" / "2026_10_ICTP" / "Absence_QuasiMajos_ICTP.pptx")

DATE = "October 30, 2026"
EQ_SRC = json.loads((EQS / "sources.json").read_text())
NUM = json.loads((FIGS / "numbers.json").read_text())

# palette: theme (QUDYMA) + poster regions
INK, INK2, GREY = "0D0D0D", "404040", "808080"
RED, ORANGE, BLUE = "BF2633", "DE9151", "6699CC"
TOPO, SKIN, SKIN_TINT, TRIV = "C8102E", "007749", "E3F0EA", "F2F2F2"
COLORS = {"red": RED, "grn": SKIN, "gry": GREY, "blu": "23427F", "org": "E07A00", "ink": INK}
MATH_FONT = "Cambria Math"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MORPH = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
         '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
         f'<p:transition xmlns:p="{P_NS}" spd="slow"><p159:morph option="byObject"/></p:transition></mc:Choice>'
         f'<mc:Fallback><p:transition xmlns:p="{P_NS}" spd="slow"><p:fade/></p:transition></mc:Fallback>'
         '</mc:AlternateContent>')

# ---------------------------------------------------------------------------
# Text with light markup: <b> <i> <sub> <sup> and colours <red> <grn> <gry> <blu> <org>.
# Greek letters and the maths symbols Montserrat lacks are set in Cambria Math
# (lower-case Greek in italics, as variables).
# ---------------------------------------------------------------------------
TAG = re.compile(r"(</?(?:b|i|sub|sup|red|grn|gry|blu|org|ink)>)")
MATH_EXTRA = set("⌊⌉≫≪∝⇒∼≲≳⁻⇏✓✗ℏ∂∞∑√≈↔")


def is_math(ch):
    return "Ͱ" <= ch <= "Ͽ" or ch in MATH_EXTRA


def split_math(text):
    out = []
    for k, grp in itertools.groupby(text, key=is_math):
        out.append(("".join(grp), k))
    return out


def parse_runs(text):
    st = {"b": False, "i": False, "sub": False, "sup": False, "color": None}
    runs = []
    for tok in TAG.split(text):
        if not tok:
            continue
        m = re.fullmatch(r"<(/?)(\w+)>", tok)
        if m and m.group(2) in ("b", "i", "sub", "sup", *COLORS):
            closing, tag = m.groups()
            if tag in ("b", "i", "sub", "sup"):
                st[tag] = not closing
            else:
                st["color"] = None if closing else COLORS[tag]
            continue
        for seg, mth in split_math(tok):
            runs.append((seg, dict(st), mth))
    return runs


def fill_paragraph(p, text, size, color=INK, bold=False, align=None):
    if align is not None:
        p.alignment = align
    for seg, st, mth in parse_runs(text):
        r = p.add_run()
        r.text = seg
        f = r.font
        f.size = Pt(size)
        f.bold = bold or st["b"] or None
        ital = st["i"] or (mth and any(c.islower() and "Ͱ" <= c <= "Ͽ" for c in seg))
        f.italic = True if ital else None
        f.color.rgb = RGBColor.from_string(st["color"] or color)
        if mth:
            f.name = MATH_FONT
        rPr = r._r.get_or_add_rPr()
        if st["sub"]:
            rPr.set("baseline", "-25000")
        if st["sup"]:
            rPr.set("baseline", "30000")
    return p


def set_bullet(p, level=0, char="•", indent=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent * (level + 1)))))
    pPr.set("indent", str(int(-Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(parse_xml(f'<a:buFont xmlns:a="{A_NS}" typeface="Arial"/>'))
    pPr.append(parse_xml(f'<a:buChar xmlns:a="{A_NS}" char="{char}"/>'))


def textbox(slide, x, y, w, h, paras, size=18, color=INK, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, space_after=6, bullets=False, name=None, line=1.0):
    """paras: list of strings (markup) or (string, dict(size=, bold=, bullet=, color=, space=))."""
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.04))
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    for k, item in enumerate(paras):
        text, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        fill_paragraph(p, text, opt.get("size", size), opt.get("color", color), opt.get("bold", False),
                       opt.get("align", align))
        p.space_after = Pt(opt.get("space", space_after))
        p.line_spacing = line
        if opt.get("bullet", bullets):
            set_bullet(p, opt.get("level", 0))
    if name:
        sh.name = name
    return sh


def box(slide, x, y, w, h, text, size=18, border=ORANGE, fill="FFFFFF", color=INK, bold=False,
        align=PP_ALIGN.CENTER, name=None, lw=2.25, radius=0.12):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if border:
        sh.line.color.rgb = RGBColor.from_string(border)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    paras = text if isinstance(text, list) else [text]
    for k, t in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        fill_paragraph(p, t, size, color, bold, align)
    if name:
        sh.name = name
    return sh


# ---------------------------------------------------------------------------
# Pictures: SVG with PNG fallback, placed at natural size (× scale)
# ---------------------------------------------------------------------------
def svg_size_in(svg_path):
    head = svg_path.read_text(errors="ignore")[:2000]
    w = re.search(r'width="([\d.]+)(pt|px)?"', head)
    h = re.search(r'height="([\d.]+)(pt|px)?"', head)
    unit = {"pt": 72.0, "px": 96.0, None: 96.0}         # unitless = CSS px (CairoMakie)
    return float(w.group(1)) / unit[w.group(2)], float(h.group(1)) / unit[h.group(2)]


def add_svg(slide, svg, png, x, y, w=None, h=None, scale=1.0, name=None, descr=None):
    W, H = svg_size_in(svg)
    if w is not None:
        scale = w / W
    elif h is not None:
        scale = h / H
    pic = slide.shapes.add_picture(str(png), Inches(x), Inches(y), Inches(W * scale), Inches(H * scale))
    package = slide.part.package
    part = Part(package.next_partname("/ppt/media/image%d.svg"), "image/svg+xml", package, svg.read_bytes())
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.blipFill.find(qn("a:blip"))
    blip.append(parse_xml(
        f'<a:extLst xmlns:a="{A_NS}"><a:ext uri="{{96DAC541-7B7A-43D3-8B79-37D633B846F1}}">'
        f'<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
        f'xmlns:r="{R_NS}" r:embed="{rid}"/></a:ext></a:extLst>'))
    if name:
        pic.name = name
    if descr:
        pic._element.nvPicPr.cNvPr.set("descr", descr)
    return pic


def fig(slide, name, x, y, morph=False, **kw):
    return add_svg(slide, FIGS / f"{name}.svg", FIGS / f"{name}.png", x, y,
                   name=("!!" if morph else "") + name, descr=kw.pop("descr", name.replace("_", " ")), **kw)


def eq(slide, name, x, y, **kw):
    return add_svg(slide, EQS / f"{name}.svg", EQS / f"{name}.png", x, y, name=name,
                   descr="LaTeX: " + EQ_SRC[name], **kw)


def photo(slide, path, x, y, w=None, h=None, name=None, descr=None):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None,
                                   Inches(h) if h else None)
    if name:
        pic.name = name
    if descr:
        pic._element.nvPicPr.cNvPr.set("descr", descr)
    return pic


def headshot(slide, path, x, y, d, top=0.02, name=None):
    """Circular crop of the upper square of a portrait photo, diameter d inches."""
    from PIL import Image
    wpx, hpx = Image.open(path).size
    side = min(wpx, hpx)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(d), Inches(d))
    t = top * hpx
    pic.crop_left = (wpx - side) / 2 / wpx
    pic.crop_right = (wpx - side) / 2 / wpx
    pic.crop_top = t / hpx
    pic.crop_bottom = max(0.0, (hpx - side - t) / hpx)
    pic.auto_shape_type = MSO_SHAPE.OVAL
    pic.line.color.rgb = RGBColor.from_string("BFBFBF")
    pic.line.width = Pt(1.25)
    if name:
        pic.name = name
    return pic


def line(slide, x0, y0, x1, y1, color="FFFFFF", width=2.0):
    c = slide.shapes.add_connector(1, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(width)
    return c


# ---------------------------------------------------------------------------
# Slide furniture: title, citation, footer, Morph, builds, notes
# ---------------------------------------------------------------------------
def set_title(slide, text):
    ph = slide.shapes.title
    tf = ph.text_frame
    tf.paragraphs[0].text = ""
    p = tf.paragraphs[0]
    for seg, st, mth in parse_runs(text):
        r = p.add_run()
        r.text = seg
        if mth:
            r.font.name = MATH_FONT
        if st["i"] or (mth and any(c.islower() for c in seg)):
            r.font.italic = True
        if st["sub"]:
            r._r.get_or_add_rPr().set("baseline", "-25000")
    return ph


def cite(slide, lines):
    """Grey 12 pt references at the bottom centre, as in the template ('CP' bold, 'et al.' italic)."""
    sh = textbox(slide, 2.42, 6.88, 8.5, 0.5, [(l, {"align": PP_ALIGN.CENTER, "space": 0}) for l in lines],
                 size=12, color=GREY, name="References")
    return sh


def footer(slide, number):
    tree = slide.shapes._spTree
    ids = [int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))]
    nid = max(ids) + 1
    tree.append(parse_xml(
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr><p:cNvPr id="{nid}" name="Date Placeholder"/>'
        '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="dt" sz="half" idx="10"/></p:nvPr>'
        '</p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" dirty="0"/>'
        f'<a:t>{DATE}</a:t></a:r></a:p></p:txBody></p:sp>'))
    tree.append(parse_xml(
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr><p:cNvPr id="{nid + 1}" name="Slide Number Placeholder"/>'
        '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="sldNum" sz="quarter" idx="12"/></p:nvPr>'
        '</p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p>'
        '<a:fld id="{E154D62F-8945-8243-A809-B5BA60694978}" type="slidenum"><a:rPr lang="en-US"/>'
        f'<a:t>{number}</a:t></a:fld></a:p></p:txBody></p:sp>'))


def morph(slide):
    sld = slide._element
    for el in list(sld):
        if el.tag in (qn("p:transition"), "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
            sld.remove(el)
    sld.find(qn("p:clrMapOvr")).addnext(parse_xml(MORPH))


def builds(slide, steps):
    """steps: list of click steps, each a list of shapes that fade in together."""
    steps = [s if isinstance(s, (list, tuple)) else [s] for s in steps if s]
    if not steps:
        return
    ids = itertools.count(3)
    pars, bld = [], []
    for step in steps:
        c_click, c_grp = next(ids), next(ids)
        effects = []
        for k, sh in enumerate(step):
            spid = sh.shape_id
            is_text = sh._element.tag == qn("p:sp")
            e, s_, a_ = next(ids), next(ids), next(ids)
            node = "clickEffect" if k == 0 else "withEffect"
            grp = ' grpId="0"' if is_text else ""
            effects.append(
                f'<p:par><p:cTn id="{e}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold"{grp} '
                f'nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr>'
                f'<p:cTn id="{s_}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility'
                f'</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{a_}" dur="500"/><p:tgtEl>'
                f'<p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')
            if is_text:
                bld.append(f'<p:bldP spid="{spid}" grpId="0" animBg="1"/>')
        pars.append(
            f'<p:par><p:cTn id="{c_click}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
            f'<p:childTnLst><p:par><p:cTn id="{c_grp}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f'<p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    bldlst = f'<p:bldLst>{"".join(bld)}</p:bldLst>' if bld else ""
    timing = parse_xml(
        f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
        f'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" '
        f'nodeType="mainSeq"><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn><p:prevCondLst>'
        f'<p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst>'
        f'<p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
        f'</p:childTnLst></p:cTn></p:par></p:tnLst>{bldlst}</p:timing>')
    sld = slide._element
    for el in sld.findall(qn("p:timing")):
        sld.remove(el)
    anchor = sld.find("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
    if anchor is None:
        anchor = sld.find(qn("p:clrMapOvr"))
    anchor.addnext(timing)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------
# Deck
# ---------------------------------------------------------------------------
class Deck:
    def __init__(self):
        self.prs = Presentation(str(TEMPLATE))
        self.layout = {l.name: l for l in self.prs.slide_layouts}
        self.number = 1
        self.order = []

    def new(self, title, number=True):
        s = self.prs.slides.add_slide(self.layout["Title Only"])
        self.number += 1
        set_title(s, title)
        if number:
            footer(s, self.number)
        morph(s)
        self.order.append(s)
        return s


def set_theme_fonts(prs, font="Montserrat"):
    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    xml = part.blob.decode("utf-8")
    xml = re.sub(r'(<a:majorFont><a:latin typeface=")[^"]*(")', rf"\g<1>{font}\2", xml)
    xml = re.sub(r'(<a:minorFont><a:latin typeface=")[^"]*(")', rf"\g<1>{font}\2", xml)
    part._blob = xml.encode("utf-8")


def drop_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)


def set_runs(shape, text, size=None):
    """Replace the text of a placeholder with markup runs, keeping its paragraph formatting."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    for r in list(p.runs):
        p._p.remove(r._r)
    for extra in tf.paragraphs[1:]:
        tf._txBody.remove(extra._p)
    for seg, st, mth in parse_runs(text):
        r = p.add_run()
        r.text = seg
        if size:
            r.font.size = Pt(size)
        if st["sup"]:
            r._r.get_or_add_rPr().set("baseline", "30000")
        if st["b"]:
            r.font.bold = True
        if mth:
            r.font.name = MATH_FONT


# ===========================================================================
# Slides
# ===========================================================================
PAPER = "<b>CP</b>, C. Robles, P. San-Jose, E. Prada, <i>Phys. Rev. Lett.</i> <b>137</b>, 096606 (2026)"
PAPER_SHORT = "<b>CP</b> <i>et al.</i> PRL 137, 096606 (2026)"


def slide_title(deck, s):
    sh = {x.name: x for x in s.shapes}
    set_runs(sh["Text Placeholder 3"],
             "ICTP–WE Heraeus College, Trieste – October 30<sup>th</sup> 2026")
    textbox(s, 1.67, 5.12, 10.0, 0.32,
            [("<b>C. Payá</b>*, C. Robles*, P. San-Jose, E. Prada · <i>Phys. Rev. Lett.</i> <b>137</b>, 096606 (2026)",
              {"align": PP_ALIGN.CENTER})], size=14, color=GREY, name="Paper")
    notes(s, "(0:20) Thank the organizers. Today: why tunneling spectroscopy in full-shell nanowires cannot be "
             "fooled by quasi-Majoranas, the most common impostor of a Majorana zero mode.")


def slide_who(deck):
    s = deck.new("Who am I?")
    W = 6.55
    bld = photo(s, ASSETS / "icmm_building.jpeg", 0.92, 1.32, w=W, name="ICMM building",
                descr="Instituto de Ciencia de Materiales de Madrid")
    h = W * 779 / 1310
    mx, my, mw = 0.92 + W - 2.05, 1.40, 1.95
    photo(s, ASSETS / "map_spain.png", mx, my, w=mw, name="Map", descr="Map of Spain")
    cx, cy, d = mx + 0.4565 * mw, my + 0.4374 * mw * 695 / 942, 0.16       # Madrid on the map
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    dot.fill.solid()
    dot.fill.fore_color.rgb = RGBColor.from_string(RED)
    dot.line.color.rgb = RGBColor.from_string("FFFFFF")
    dot.line.width = Pt(1.5)
    dot.shadow.inherit = False
    dot.name = "Madrid"
    lb = box(s, cx + 0.12, cy - 0.15, 0.95, 0.3, "<b>Madrid</b>", size=12, border=None, fill="FFFFFF", radius=0.3)
    lb.text_frame.margin_left = lb.text_frame.margin_right = Inches(0.04)
    textbox(s, 0.92, 1.32 + h + 0.06, W, 0.32,
            ["Instituto de Ciencia de Materiales de Madrid (ICMM), CSIC"], size=14, color=INK2)
    textbox(s, 0.92, 1.32 + h + 0.48, W, 1.2, [
        ("<b>Carlos Payá</b> · PhD candidate · defence <red><b>March 2027</b></red>", {"size": 18}),
        ("Thesis: full-shell hybrid nanowires, from Caroli–de Gennes–Matricon analogs to Majoranas "
         "and Josephson junctions", {"size": 14, "color": INK2}),
    ])
    # supervisors
    x0 = 8.0
    textbox(s, x0, 1.32, 4.4, 0.4, ["<b>PhD supervisors</b>"], size=18)
    for k, (f, n) in enumerate((("face_elsa.png", "Elsa Prada"), ("face_ramon.jpg", "Ramón Aguado"))):
        x = x0 + 0.25 + k * 2.15
        headshot(s, ASSETS / f, x, 1.82, 1.45, name=n)
        textbox(s, x - 0.3, 3.32, 2.05, 0.32, [n], size=14, align=PP_ALIGN.CENTER)
    # this work
    textbox(s, x0, 3.92, 4.4, 0.4, ["<b>This work</b>"], size=18)
    for k, (f, n, role) in enumerate((("face_cesar.png", "César Robles*", "MSc student"),
                                      ("face_pablo.png", "Pablo San-Jose", "ICMM"),
                                      ("face_elsa.png", "Elsa Prada", "ICMM"))):
        x = x0 + 0.1 + k * 1.45
        headshot(s, ASSETS / f, x, 4.42, 1.15, name=n)
        textbox(s, x - 0.17, 5.62, 1.5, 0.55, [(n, {"space": 0}), (role, {"size": 12, "color": GREY})],
                size=13, align=PP_ALIGN.CENTER)
    textbox(s, x0, 6.28, 4.45, 0.3, ["*equal contribution"], size=12, color=GREY)
    notes(s, "(1:00) I am a PhD candidate at ICMM-CSIC in Madrid, supervised by Elsa Prada and Ramón Aguado, "
             "and I will defend in March 2027. My thesis is on full-shell hybrid nanowires: the Caroli–de "
             "Gennes–Matricon analogs they host, their Majoranas and their Josephson junctions. This work was "
             "done with César Robles, a master student who shares first authorship, Pablo San-Jose and Elsa Prada.")


def slide_majo_tunneling(deck):
    s = deck.new("Majoranas in tunneling spectroscopy")
    sk = photo(s, ASSETS / "sketch_partial_3d_cut_hi.png", 0.92, 1.42, w=5.5, name="!!sketch_partial",
               descr="Partial-shell nanowire with a tunnel probe")
    t1 = textbox(s, 0.92, 3.62, 5.9, 0.4, ["Topological phase for"], size=18)
    e1 = eq(s, "eq_vzc", 1.25, 4.02)
    t2 = textbox(s, 0.92, 4.72, 5.9, 0.9, [
        ("Majorana zero modes (MZMs) at the wire ends", {"bullet": True}),
        ("Local probe: <i>dI</i>/<i>dV</i> ∝ LDOS at the end → zero-bias peak (ZBP)", {"bullet": True}),
    ], size=16)
    pd = fig(s, "s03_pd_partial", 7.42, 1.30, morph=True, descr="Phase diagram of the partial-shell wire")
    ld = fig(s, "s03_ldos_vB_sharp", 7.42, 4.40, morph=True,
             descr="End LDOS versus Zeeman field for an abrupt wire end")
    cb = box(s, 0.92, 5.78, 5.9, 0.82, "A ZBP is <b>necessary but not sufficient</b>: trivial states can also "
             "stick to zero energy", size=16)
    builds(s, [[t1, e1, pd], [t2, ld], [cb]])
    cite(s, ["R. M. Lutchyn <i>et al.</i> PRL 105, 077001 (2010) | Y. Oreg <i>et al.</i> PRL 105, 177002 (2010)",
             "V. Mourik <i>et al.</i> Science 336, 1003 (2012) | E. Prada <i>et al.</i> Nat. Rev. Phys. 2, 575 (2020)"])
    notes(s, "(1:00) Quick recap of the Majorana nanowire you saw in the school. A semiconductor with strong "
             "spin–orbit coupling, proximitized by a superconductor on some facets, under a Zeeman field. It is "
             "topological for V_Z above V_Z^c = sqrt(Δ0² + μ²): the red wedge. A tunnel probe at the end measures "
             "the local density of states. For a clean wire with a sharp end, the bulk gap closes at V_Z^c and a "
             "zero-bias peak appears after it: the textbook signature (illustrative calculation, same model as "
             "the paper with χ = 1 nm). But a ZBP alone is not proof: trivial states can also stick to zero.")


def slide_smooth_partial(deck):
    s = deck.new("Smooth confinement at the wire's end")
    cut = fig(s, "s04_cut_partial", 0.92, 1.34, descr="Longitudinal cut of the partial-shell wire")
    # electrostatics only, with the first text; the MZM strip fades in on top with the MZM path
    bare = fig(s, "s04_rs_partial_bare", 0.92, 2.12,
               descr="Electrostatics at the wire's end: insulating barrier and hybrid wire")
    rm = fig(s, "s04_rs_partial_majo", 0.92, 2.12, descr="Local phases and MZM wavefunction, V_Z^(1)")
    rq = fig(s, "s04_rs_partial_qmajo", 0.92, 4.02, morph=True,
             descr="Local phases and quasi-MZM wavefunctions, V_Z^(2)")
    pd = fig(s, "s03_pd_partial", 7.42, 1.30, morph=True, descr="Phase diagram of the partial-shell wire")
    pm = fig(s, "s04_pd_partial_mzm", 7.42, 1.30, descr="mu(z) path for V_Z^(1)")
    pq = fig(s, "s04_pd_partial_qmzm", 7.42, 1.30, descr="mu(z) path for V_Z^(2)")
    t1 = textbox(s, 7.42, 4.48, 5.0, 0.62, ["The tunnel gate depletes the end; the potential heals over a length "
                                            "<i>χ</i>:"], size=16)
    e1 = eq(s, "eq_muz", 7.85, 5.10)
    t2 = textbox(s, 7.42, 5.65, 5.0, 0.4, ["Locally topological wherever"], size=16)
    e2 = eq(s, "eq_local", 7.85, 6.03)
    t3 = textbox(s, 0.92, 5.95, 6.3, 0.62, ["Below <i>V</i><sub>Z</sub><sup>c</sup>, only a segment of length ~<i>χ</i> "
                                            "turns topological: <b>two boundaries, two Majorana components</b>"],
                 size=16)
    builds(s, [[t1, e1, bare], [t2, e2], [rm, pm], [rq, pq, t3]])
    cite(s, [PAPER_SHORT,
             "E. Prada <i>et al.</i> PRB 86, 180503(R) (2012) | G. Kells <i>et al.</i> PRB 86, 100503(R) (2012)"])
    notes(s, "(1:15) In a real device the tunnel gate depletes the end of the wire, and the potential heals "
             "smoothly over a length χ: μ(z) = μ_bulk(1 − e^{−z/χ}). Each point of the wire sits somewhere in the "
             "phase diagram, so moving along the wire traces a vertical path. Above V_Z^c the path goes from "
             "insulating straight into the topological phase: the whole wire is topological and the MZM sits "
             "at the end. Below V_Z^c, at V_Z^(2), the path enters the topological region and leaves it again: "
             "only a segment of length ~χ is topological. That segment has two boundaries, and each hosts a "
             "Majorana component.")


def slide_quasimajo(deck):
    s = deck.new("Quasi-Majoranas: trivial zero modes")
    rq = fig(s, "s04_rs_partial_qmajo", 0.92, 1.34, morph=True,
             descr="Local phases and quasi-MZM wavefunctions, V_Z^(2)")
    t1 = textbox(s, 0.92, 3.30, 6.3, 1.65, [
        ("Two overlapping Majorana components, the <b>quasi-MZMs</b>", {"size": 18}),
        ("<blu>|Ψ<sub>L</sub>|<sup>2</sup></blu> at the abrupt end: Fermi oscillations", {"bullet": True}),
        ("<org>|Ψ<sub>R</sub>|<sup>2</sup></org> at the smooth interface: no oscillations", {"bullet": True}),
        ("→ no coherent splitting: <b>exponentially pinned to zero energy</b> when", {"size": 16}),
    ], size=16, space_after=4)
    e1 = eq(s, "eq_chi", 2.4, 5.02)
    dos = fig(s, "s05_dos_partial", 7.42, 1.30, descr="DOS where mu(z) varies versus Zeeman field")
    t2 = textbox(s, 7.42, 4.55, 5.0, 1.4, [
        ("The bulk gap closes at <i>V</i><sub>Z</sub><sup>c</sup>: MZM above it", {"bullet": True}),
        ("The quasi-MZM peak sits at zero <b>before</b> <i>V</i><sub>Z</sub><sup>c</sup>, in the trivial phase",
         {"bullet": True}),
    ], size=16)
    builds(s, [[t1, e1], [dos, t2]])
    cite(s, [PAPER_SHORT,
             "G. Kells <i>et al.</i> PRB 86, 100503(R) (2012) | F. Peñaranda <i>et al.</i> PRB 98, 235406 (2018)"])
    notes(s, "(1:15) These are quasi-Majoranas. The outer component is abruptly confined at the end and "
             "oscillates with the Fermi momentum; the inner one sits at a smooth interface and has no "
             "oscillations. Although they overlap strongly, they cannot hybridize coherently, so their splitting "
             "is exponentially small in the smoothness (Kells, Meidan, Brouwer 2012): they pin to zero energy for "
             "χ ≫ ħαμ_bulk/(V_Z Δ0). In the DOS of the region where μ(z) varies we see the bulk transition at "
             "V_Z^c and the MZM after it; but the quasi-MZM zero-energy peak appears well before V_Z^c, in the "
             "trivial phase.")


def slide_false_positive(deck):
    s = deck.new("False positives in local spectroscopy")
    a = fig(s, "s03_ldos_vB_sharp", 0.92, 1.30, descr="End LDOS versus Zeeman field, sharp end")
    b = fig(s, "s06_ldos_vB_smooth", 0.92, 3.86, descr="End LDOS versus Zeeman field, smooth end")
    cm = fig(s, "s06_ldos_chi_partial_majo", 6.82, 1.18, descr="End LDOS versus smoothness, MZM")
    cq = fig(s, "s06_ldos_chi_partial_qmajo", 6.82, 1.18, descr="End LDOS versus smoothness, quasi-MZM")
    tg = fig(s, "s06_ldos_chi_partial_tag", 6.82, 1.18, descr="False positive tag")
    t1 = textbox(s, 0.92, 6.30, 5.6, 0.45, ["The ZBP starts <b>before</b> the bulk transition"], size=16)
    q = box(s, 6.95, 6.08, 5.45, 0.72, "Can a different geometry defuse this false positive?", size=16,
            border=SKIN, bold=True)
    builds(s, [[b, t1], [cm], [cq], [tg], [q]])
    cite(s, [PAPER_SHORT, "A. Vuik <i>et al.</i> SciPost Phys. 7, 061 (2019) | C.-X. Liu <i>et al.</i> PRB 96, "
                          "075161 (2017)"])
    notes(s, "(1:15) This is what the probe actually sees. With a sharp end the ZBP appears only after V_Z^c; "
             "with a smooth end (χ = 200 nm, published data) it starts well before: the probe cannot see the bulk "
             "transition, and the ZBP looks the same on both sides. As a function of smoothness: the true MZM "
             "gives a ZBP for any χ; the quasi-MZM splits for sharp ends but pins to zero above ~100–200 nm. "
             "Same zero-bias peak, no topology: a false positive. Quasi-MZMs also reproduce quantized 2e²/h peaks "
             "and 4π Josephson effects (Vuik et al. 2019). So the question is whether a different geometry can "
             "defuse it.")


def slide_fullshell(deck):
    s = deck.new("Full-shell hybrid nanowires")
    sk = photo(s, ASSETS / "sketch_full_3d_cut_hi.png", 0.92, 1.42, w=5.5, name="!!sketch_full",
               descr="Full-shell nanowire with a tunnel probe")
    t1 = textbox(s, 0.92, 3.50, 5.7, 1.0, [
        ("Superconducting shell fully wraps the semiconductor core", {"bullet": True}),
        ("Axial field → flux through the shell:", {"bullet": True}),
    ], size=16)
    e1 = eq(s, "eq_flux", 1.25, 4.35)
    t2 = textbox(s, 0.92, 5.12, 5.7, 1.5, [
        ("Fluxoid <i>n</i>: <b>Little–Parks lobes</b> of the shell gap", {"bullet": True}),
        ("Topological transition driven by the flux (orbital effect + radial SOC) at <i>B</i> ≈ 0.1 T, "
         "in <b>odd lobes</b>", {"bullet": True}),
    ], size=16)
    img = photo(s, ASSETS / "vaitiekenas2020_fig1.jpeg", 6.82, 1.42, w=5.6, name="Vaitiekenas 2020",
                descr="Full-shell device and tunneling spectroscopy, Vaitiekenas et al. Science 2020")
    cap = textbox(s, 6.82, 4.55, 5.6, 0.3, ["S. Vaitiekėnas <i>et al.</i>, Science 367, eaav3392 (2020)"],
                  size=12, color=GREY, align=PP_ALIGN.RIGHT)
    t3 = textbox(s, 6.82, 4.95, 5.6, 0.45, [("Zero-bias peaks observed in the first lobe", {"bullet": True})],
                 size=16)
    t4 = textbox(s, 6.82, 5.42, 5.6, 1.0, [("…but trivial ZBPs also appear, from flux-tunable Andreev states in "
                                            "the junction", {"bullet": True})], size=16)
    builds(s, [[t2], [img, cap, t3], [t4]])
    cite(s, ["S. Vaitiekėnas <i>et al.</i> Science 367, eaav3392 (2020) | <b>CP</b> <i>et al.</i> PRB 109, "
             "115428 (2024)", "M. Valentini <i>et al.</i> Science 373, 82 (2021)"])
    notes(s, "(1:00) The alternative: full-shell nanowires, where the superconductor wraps the whole core. An "
             "axial field threads a flux through the shell and the order parameter winds n times around the "
             "wire, with n the fluxoid. This gives Little–Parks lobes in the shell gap. The topological "
             "transition is now driven by the flux, through the orbital effect and the radial spin–orbit "
             "coupling, at fields around 0.1 T, and only in odd lobes. Vaitiekėnas and coworkers saw zero-bias "
             "peaks in the first lobe in 2020.")


def slide_cdgm(deck):
    s = deck.new("A synthetic vortex: CdGM analogs")
    v = fig(s, "s08_cartoon_vortex", 0.92, 1.38, descr="Abrikosov vortex cartoon")
    tv = textbox(s, 3.62, 1.50, 2.85, 1.4, [
        ("<b>Abrikosov vortex</b>", {"size": 16}),
        ("Andreev reflection off the <i>Δ</i>(<i>r</i>) wall confines Caroli–de Gennes–Matricon states",
         {"size": 14}),
    ], size=14)
    ev = eq(s, "eq_cdgm_vortex", 3.68, 2.95, scale=0.82)
    f = fig(s, "s08_cartoon_fullshell", 0.92, 4.05, descr="Full-shell wire cartoon")
    tf = textbox(s, 3.62, 4.17, 2.85, 1.5, [
        ("<b>Full-shell wire</b>", {"size": 16}),
        ("The winding shell confines Andreev subbands in the core: <b>CdGM analogs</b>, one per "
         "<i>m</i><sub>J</sub>", {"size": 14}),
    ], size=14)
    ef = eq(s, "eq_cdgm_fs", 3.68, 5.70, scale=0.82)
    tf2 = textbox(s, 4.55, 5.66, 1.95, 0.4, ["→ resolvable"], size=14)
    b0 = fig(s, "s08_bands_n0", 6.52, 1.25, descr="Hollow-core bands and end LDOS, flux 0")
    b1 = fig(s, "s08_bands_n1", 6.52, 1.25, descr="Hollow-core bands and end LDOS, flux 0.75")
    cap = textbox(s, 6.6, 6.28, 5.8, 0.5, ["Hollow-core model (illustrative): each subband's gap edge, a Van Hove "
                                           "singularity, gives one peak"], size=12, color=GREY)
    builds(s, [[f, tf, ef, tf2], [b0], [b1, cap]])
    cite(s, ["C. Caroli, P. G. de Gennes, J. Matricon, Phys. Lett. 9, 307 (1964)",
             "P. San-Jose, <b>CP</b> <i>et al.</i> PRB 107, 155423 (2023)"])
    notes(s, "(1:15) Why is the full shell special? A phase winding around a normal region is a vortex. In an "
             "Abrikosov vortex, the pairing vanishes in the core and Andreev reflection off that wall confines "
             "the Caroli–de Gennes–Matricon states, with tiny spacings Δ²/2E_F, about 10⁻³ of the gap in a "
             "conventional superconductor. The full shell is a synthetic "
             "vortex whose core is the semiconductor: Andreev reflection at the core–shell interface confines 1D "
             "subbands, one per generalized angular momentum m_J. Each subband has a gap edge, a Van Hove "
             "singularity, so they show up as sharp peaks with spacings comparable to the gap: the CdGM "
             "analogs, resolvable in tunneling. (Right: illustrative hollow-core calculation. At Φ = 0 all gap edges coincide; inside "
             "lobe 1 they spread into a ladder.)")


def slide_cdgm_exp(deck):
    """Two slides joined by Morph: the experiment with the shared colorbar next to it (9a), then
    the theory appears and the colorbar ("!!cbar" on both) slides to its right (9b)."""
    def common(s, cbar_x):
        photo(s, ASSETS / "deng2025_cross_section.png", 0.92, 1.38, w=2.25, name="Cross-section",
              descr="InAs/Al full-shell cross-section")
        h1 = 2.25 * 1933 / 2142
        line(s, 1.08, 1.38 + h1 - 0.16, 1.08 + 50 / 0.12631578947368421 / 2142 * 2.25, 1.38 + h1 - 0.16)
        textbox(s, 1.0, 1.38 + h1 - 0.47, 1.0, 0.28, ["50 nm"], size=12, color="FFFFFF")
        textbox(s, 0.92, 1.38 + h1 + 0.03, 2.25, 0.3, ["InAs core, 7 nm Al shell"], size=12, color=GREY,
                align=PP_ALIGN.CENTER)
        photo(s, ASSETS / "deng2025_device1.png", 1.42, 3.82, h=2.95, name="Device",
              descr="False-colour SEM of device 1")
        wsem = 2.95 * 1483 / 3458
        # scale bar at the lower right of the SEM image (Carlos, 2026-10-20)
        line(s, 2.157, 6.436, 2.157 + 1000 / 1.9108280254777081 / 1483 * wsem, 6.436)
        textbox(s, 2.156, 6.16, 0.8, 0.28, ["1 µm"], size=12, color="FFFFFF")
        fig(s, "s09_exp", 3.52, 1.28, descr="Measured dI/dV versus field and bias, device 1")
        add_svg(s, FIGS / "s09_cbar.svg", FIGS / "s09_cbar.png", cbar_x, 1.28, name="!!cbar",
                descr="Shared colour bar, dI/dV in e^2/h")
        cite(s, ["M. T. Deng, <b>CP</b>, P. San-Jose, E. Prada, C. M. Marcus, S. Vaitiekėnas, PRL 134, 206302 (2025)"])

    # colorbar picture: bar 0.02 in from its left edge; panels at 0.85–4.30 (exp), 4.45–7.90 (theory)
    s = deck.new("CdGM analogs in the experiment")
    common(s, 3.52 + 4.45 - 0.02)
    notes(s, "(0:25) And we see them. Tunneling spectroscopy on InAs/Al full-shell wires from the Copenhagen "
             "group: the zeroth lobe has a hard gap, while the first and second lobes are filled with discrete "
             "subgap states that disperse and skew with flux.")
    s = deck.new("CdGM analogs in the experiment")
    common(s, 3.52 + 8.05 - 0.02)
    fig(s, "s09_theory", 3.52, 1.28, descr="Modelled dI/dV versus flux and bias")
    t1 = textbox(s, 3.62, 5.85, 8.8, 0.9, [
        ("Lobe 0: hard gap. Lobes 1, 2: discrete subgap states that disperse and skew with flux, one per "
         "<i>m</i><sub>J</sub> subband", {"bullet": True}),
    ], size=16)
    builds(s, [[t1]])
    notes(s, "(0:20) Our model reproduces them, including the skewness, which tells us the average radius and "
             "angular momentum of each subband. So a full-shell wire never has an empty odd lobe: CdGM analogs "
             "are always there.")


def slide_mj(deck):
    s = deck.new("Only <i>m</i><sub>J</sub> = 0 can host Majoranas")
    tot = fig(s, "s10_dos_lobes_total", 6.12, 1.22, descr="DOS versus flux, all angular sectors")
    mj0 = fig(s, "s10_dos_lobes_mj0", 6.12, 1.22, descr="DOS versus flux, m_J = 0 sector")
    t1 = textbox(s, 0.92, 1.36, 5.1, 0.4, ["Cylindrical symmetry:"], size=16)
    e1 = eq(s, "eq_jz", 1.05, 1.78, scale=0.85)
    t1b = textbox(s, 0.92, 2.22, 5.1, 0.7, ["→ one Hamiltonian per sector <i>m</i><sub>J</sub>; CdGM analogs "
                                            "live in all of them"], size=16)
    t2 = textbox(s, 0.92, 3.00, 5.1, 0.4, ["Particle–hole symmetry maps <i>m</i><sub>J</sub> → −<i>m</i><sub>J</sub>:"],
                 size=16)
    e2 = eq(s, "eq_phs", 1.05, 3.42, scale=0.9)
    t3 = textbox(s, 0.92, 3.92, 5.1, 0.4, ["An unpaired MZM needs <i>m</i><sub>J</sub> = −<i>m</i><sub>J</sub>:"],
                 size=16)
    e3 = eq(s, "eq_mj0", 1.05, 4.35, scale=0.9)
    t3b = textbox(s, 3.15, 4.30, 2.9, 0.4, ["→ only <b>odd lobes</b>"], size=16)
    cb = box(s, 0.92, 5.20, 5.05, 1.0, "MZMs typically coexist with CdGM analogs of other <i>m</i><sub>J</sub> at "
             "zero energy: <b>gapless</b>", size=16, border=ORANGE)
    builds(s, [[t1, e1, t1b], [t2, e2, t3, e3, t3b], [mj0], [cb]])
    cite(s, ["<b>CP</b> <i>et al.</i> PRB 109, 115428 (2024) | F. Peñaranda <i>et al.</i> PRR 2, 023171 (2020)"])
    notes(s, "(1:15) The DOS of our device shows the Little–Parks lobes filled with CdGM analogs. In a "
             "cylindrical approximation the generalized angular momentum J_z is conserved, so the Hamiltonian "
             "splits into sectors m_J. Particle–hole symmetry maps m_J to −m_J, so an unpaired Majorana needs the "
             "self-conjugate sector m_J = 0, which exists only for odd fluxoid. The bottom panel shows only that "
             "sector: it exists in lobe 1 and not in lobes 0 and 2. And the MZM has to share zero energy with the "
             "CdGM analogs of all other sectors: the topological phase is generically gapless.")


def slide_smooth_full(deck):
    s = deck.new("Smooth confinement in full shells")
    cut = fig(s, "s11_cut_full", 0.92, 1.40, descr="Longitudinal cut of the full-shell wire")
    t1 = textbox(s, 0.92, 2.22, 6.0, 0.7, ["Laplace equation inside the grounded shell, with the gate depleting "
                                           "the end (<i>z</i> = 0):"], size=16)
    e1 = eq(s, "eq_bessel", 0.98, 2.98, scale=0.92)
    t2 = textbox(s, 0.92, 3.95, 6.0, 0.65, ["The slowest mode sets the smoothness <i>χ</i> (<i>R</i>/<i>k</i><sub>1</sub> "
                                            "without screening; free parameter)"], size=14, color=INK2)
    pd = fig(s, "s11_pd_full", 7.02, 1.22, descr="Phase diagram of the full-shell wire versus flux and mu")
    pm = fig(s, "s11_pd_full_mzm", 7.02, 1.22, descr="mu(z) path for Phi^(1)")
    pq = fig(s, "s11_pd_full_qmzm", 7.02, 1.22, descr="mu(z) path for Phi^(2)")
    ps = fig(s, "s11_pd_full_skin", 7.02, 1.22, descr="Trivial skin label")
    cb = box(s, 0.92, 4.85, 6.0, 1.05, "<i>μ</i>(<i>z</i>) must cross a trivial region before turning "
             "topological: a <grn><b>trivial skin</b></grn> at the wire's end", size=16, border=SKIN)
    t3 = textbox(s, 7.12, 5.35, 5.3, 0.75, ["MZM for <i>Φ</i> < <i>Φ</i><sup>c</sup>, quasi-MZM for "
                                            "<i>Φ</i> > <i>Φ</i><sup>c</sup>: inverted with respect to partial "
                                            "shells"], size=14, color=INK2)
    builds(s, [[t2], [pm], [pq], [ps, cb], [t3]])
    cite(s, [PAPER_SHORT + ", End Matter A"])
    notes(s, "(1:15) Now the same smooth confinement in a full shell. The grounded shell screens external gates, "
             "so their potential only enters through the uncovered end; solving Laplace's equation inside the "
             "core gives a sum of evanescent Bessel modes, and the slowest one sets the smoothness χ, which we "
             "keep as a free parameter. Again μ(z) traces a vertical path in the phase diagram, now versus flux "
             "inside lobe 1. The MZM case: the path does not enter the topological region directly, it first "
             "crosses a trivial region. Same for the quasi-MZM case. That trivial segment at the end of the wire "
             "is the trivial skin. Note the inversion: here the MZM is at the lower flux and the quasi-MZM at the "
             "higher one.")


def slide_skin_origin(deck):
    s = deck.new("Origin of the trivial skin")
    fb = fig(s, "s12_filling_bands", 0.92, 1.32, descr="Normal-state angular subbands at k_z = 0")
    fs = fig(s, "s12_filling_skin", 0.92, 1.32, descr="Threshold mu^ts and trivial skin")
    pd = fig(s, "s12_pd_mualpha", 6.82, 1.32, descr="Phase diagram versus mu and average SOC")
    t1 = textbox(s, 0.92, 5.45, 5.55, 1.3, ["Radial SOC moves the band bottom away from the <i>m</i><sub>J</sub> = 0 "
                                            "modes: CdGM-analog subbands fill first, so topology needs "
                                            "<i>μ</i> > <grn><i>μ</i><sup>ts</sup></grn>"], size=16)
    t2 = textbox(s, 6.95, 5.45, 5.45, 1.3, ["Partial shells have <b>no threshold</b>: the lowest subband turns "
                                            "topological right above depletion, so there is no skin"], size=16)
    builds(s, [[fs, t1], [pd], [t2]])
    cite(s, [PAPER_SHORT + ", End Matter B and Fig. 4"])
    notes(s, "(1:15) Where does the skin come from? Look at the normal-state subbands at k_z = 0 versus angular "
             "momentum. The radial spin–orbit coupling shifts the parabolas, so the band bottom is not at the "
             "m_J = 0 modes, the only ones that can become topological. As we fill the wire, the CdGM-analog "
             "subbands are populated first and the m_J = 0 modes only start at a threshold μ^ts. Realistic SOC "
             "makes the first radial mode useless, so in our wire the threshold comes from the second radial "
             "mode, at about 22 meV: μ(z) has to climb all the way there through a trivial region. In a partial "
             "shell there is no such threshold.")


def slide_deep(deck):
    s = deck.new("Quasi-Majoranas form, but deep inside")
    cut = fig(s, "s11_cut_full", 0.92, 1.34, descr="Longitudinal cut of the full-shell wire")
    rm = fig(s, "s13_rs_full_majo", 0.92, 2.12, descr="Local phases and MZM wavefunction, Phi^(1)")
    rq = fig(s, "s13_rs_full_qmajo", 0.92, 4.02, descr="Local phases and quasi-MZM wavefunctions, Phi^(2)")
    t1 = textbox(s, 0.92, 5.95, 6.3, 0.62, [f"A trivial skin ≈ {NUM['full_skin_end_chi_Phi2']:.1f}<i>χ</i> wide "
                                            "separates the probe from the MZM <b>and</b> the quasi-MZMs"], size=16)
    dos = fig(s, "s13_dos_mj0", 7.32, 1.22, descr="m_J = 0 DOS versus flux in the first lobe")
    t2 = textbox(s, 7.42, 5.05, 5.0, 1.2, ["Quasi-MZMs do form: zero-energy peak for <i>Φ</i> > <i>Φ</i><sup>c</sup>, "
                                           "with exponentially small splitting"], size=16)
    builds(s, [[rq, t1], [dos, t2]])
    cite(s, [PAPER_SHORT + ", Fig. 1"])
    notes(s, "(1:15) In real space: at Φ^(1) the wire is topological beyond the skin, and the MZM sits at the end "
             "of the skin, about 3.4 χ from the probe. At Φ^(2) the topological segment is now between the skin "
             "and the bulk, and the two quasi-MZMs form at its boundaries. They are real: the m_J = 0 DOS shows "
             "the MZM peak below Φ^c and the quasi-MZM peak above it, pinned to zero. So quasi-Majoranas form in "
             "full shells too, but behind the skin.")


def slide_no_fp(deck):
    """Two slides joined by Morph. 15a: partial row; on click the full-shell true-MZM panel with
    its colorbar right next to it. 15b: the quasi-MZM panel arrives and the colorbar ("!!cbarF"
    on both) slides to the right of it; then annotations, tag and box."""
    # grid canvas at x = 0.92: true-MZM panel 2.35–6.35, quasi-MZM 6.75–10.75, colorbars at 10.92;
    # the colorbar picture has its bar 0.02 in from its left edge
    cb_near, cb_far = 0.92 + (2.35 + 4.0 + 0.17) - 0.02, 0.92 + 10.92 - 0.02

    def common(s):
        lay = {n: fig(s, n, 0.92, 1.25, descr=n.replace("_", " "))
               for n in ("s14_grid_partial", "s14_grid_tagP", "s14_grid_fullM")}
        cite(s, [PAPER_SHORT + ", Fig. 2"])
        return lay

    s = deck.new("No false positives in full shells")
    lay = common(s)
    cb = add_svg(s, FIGS / "s14_grid_cbarF.svg", FIGS / "s14_grid_cbarF.png", cb_near, 1.25, name="!!cbarF",
                 descr="Full-shell row colour bar, LDOS (arb. u.)")
    builds(s, [[lay["s14_grid_fullM"], cb]])
    notes(s, f"(0:50) The key result: end LDOS versus smoothness. Partial shell, as before: the MZM gives a ZBP "
             f"for any χ and the quasi-MZM becomes indistinguishable for smooth ends: a false positive. Full "
             f"shell, true MZM: visible for sharp ends, but the skin grows with χ and the end signal is gone "
             f"beyond ~{NUM['chi_vanish_full_MZM_nm']:.0f} nm.")

    s = deck.new("No false positives in full shells")
    common(s)
    fig(s, "s14_grid_fullQ", 0.92, 1.25, descr="s14 grid fullQ")
    add_svg(s, FIGS / "s14_grid_cbarF.svg", FIGS / "s14_grid_cbarF.png", cb_far, 1.25, name="!!cbarF",
            descr="Full-shell row colour bar, LDOS (arb. u.)")
    ann = fig(s, "s14_grid_ann", 0.92, 1.25, descr="s14 grid ann")
    tagF = fig(s, "s14_grid_tagF", 0.92, 1.25, descr="s14 grid tagF")
    box_ = box(s, 0.92, 6.08, 11.5, 0.72, "The smooth potential that creates the impostor also hides it: <b>a ZBP at "
               "the end of a full-shell wire cannot be a quasi-Majorana</b>", size=17, border=SKIN)
    builds(s, [[ann], [tagF], [box_]])
    notes(s, "(0:55) Quasi-MZM: some split states for very sharp ends, and then nothing. Nothing appears at the "
             "end, not even when the quasi-MZMs pin to zero energy deep inside the wire. The same smooth "
             "potential that creates the impostor hides it. Caveat: the skin also hides true MZMs when the end "
             "is smooth.")


def slide_robust(deck):
    s = deck.new("Robust to disorder and temperature")
    dp = fig(s, "s15_disorder_partial", 0.92, 1.32, descr="End LDOS with short-range disorder, partial shell")
    df = fig(s, "s15_disorder_full", 0.92, 1.32, descr="End LDOS with short-range disorder, full shell")
    t1 = textbox(s, 7.75, 1.55, 4.65, 3.3, [
        ("<b>Short-range disorder</b>, weaker than the topological minigap", {"size": 16, "space": 10}),
        ("MZM zero-bias peaks survive in both geometries", {"bullet": True}),
        ("Quasi-MZM peaks smear out in partial shells, and never reach the end of full shells",
         {"bullet": True}),
    ], size=16)
    t2 = textbox(s, 7.75, 4.65, 4.65, 1.4, [
        ("Confirmed by the tunneling conductance, including <b>all <i>m</i><sub>J</sub> sectors</b> and "
         "finite temperature", {"size": 16}),
    ], size=16)
    builds(s, [[df], [t1], [t2]])
    cite(s, [PAPER_SHORT + ", Fig. 3 and Supplemental Material S-IV, S-V"])
    notes(s, "(0:50) Two checks. With short-range disorder below the minigap, true MZM peaks survive in both "
             "geometries, the quasi-MZM peak gets smeared in the partial shell, and in the full shell nothing "
             "reaches the end. And the tunneling conductance including all angular sectors and finite "
             "temperatures gives the same picture (Supplemental Material).")


def slide_conclusions(deck):
    s = deck.new("Conclusions & outlook")
    rows = [("s16_thumb_partial", "<b>Partial shells</b>: smooth confinement creates quasi-MZMs that the local "
                                  "probe cannot tell from MZMs"),
            ("s16_thumb_skin", "<b>Full shells</b>: CdGM analogs fill first, so <i>μ</i>(<i>z</i>) crosses a "
                               "<grn><b>trivial skin</b></grn> before turning topological"),
            ("s16_thumb_full", "The skin hides the quasi-MZMs: <b>no false positives</b> from smooth confinement")]
    for k, (f, t) in enumerate(rows):
        y = 1.38 + k * 1.52
        fig(s, f, 0.92, y, descr=f.replace("_", " "))
        textbox(s, 3.45, y + 0.05, 3.25, 1.35, [t], size=15, anchor=MSO_ANCHOR.MIDDLE)
    out = textbox(s, 7.05, 1.38, 5.35, 3.3, [
        ("<b>Outlook</b>", {"size": 20, "color": RED, "space": 8}),
        ("<b>If you see a ZBP</b>: it cannot be a quasi-MZM. Other trivial origins (e.g. flux-tunable Andreev "
         "states) must still be excluded", {"bullet": True, "space": 8}),
        ("<b>If you don't</b>: the skin buries true MZMs too. Sharp barriers give the most informative "
         "measurements", {"bullet": True, "space": 8}),
        ("Beyond tunneling: CdGM analogs also leave fingerprints in the Josephson critical current "
         "(<b>CP</b> <i>et al.</i> arXiv:2607.18371)", {"bullet": True}),
    ], size=15)
    key = box(s, 0.92, 6.0, 11.5, 0.8, "In full-shell nanowires, the smooth potential that creates Majorana "
              "impostors also hides them: <b>zero-bias peaks cannot be quasi-Majoranas</b>", size=17,
              border=ORANGE)
    builds(s, [[out], [key]])
    cite(s, [PAPER.replace("<i>Phys. Rev. Lett.</i>", "PRL")])
    notes(s, "(1:30) To conclude. In partial shells, smooth confinement creates quasi-Majoranas that local "
             "spectroscopy cannot tell from MZMs. Full shells are synthetic vortices: CdGM analogs fill first, "
             "so μ(z) has to cross a trivial skin before reaching the topological phase, and that skin hides the "
             "quasi-Majoranas. If you see a zero-bias peak in a full shell, it is not a quasi-Majorana, although "
             "other mechanisms such as flux-tunable Andreev states in the junction still have to be excluded. "
             "If you don't see one, the skin may be hiding a true MZM too: sharp barriers are the way to go. "
             "And beyond tunneling, the CdGM analogs also show up in the critical current of full-shell "
             "Josephson junctions, our latest work with the Madrid experimental group. Thank you.")


def slide_thanks(deck, s):
    sh = {x.name: x for x in s.shapes}
    set_runs(sh["Text Placeholder 3"], "ICTP Trieste – October 30<sup>th</sup> 2026")
    set_runs(sh["Text Placeholder 6"], "PID2021-125343NB-I00 · PID2023-150224NB-I00", size=11)
    sh["Text Placeholder 6"].top, sh["Text Placeholder 6"].height = Inches(5.08), Inches(0.5)
    p2 = sh["Text Placeholder 6"].text_frame.add_paragraph()
    fill_paragraph(p2, "PRE2022-101362 · CEX2024-001445-S", 11, color="A6A6A6")
    for r in sh["Text Placeholder 6"].text_frame.paragraphs[0].runs:
        r.font.color.rgb = RGBColor.from_string("A6A6A6")

    def shrink_names(shapes):
        for x in shapes:
            if x.shape_type == 6:
                shrink_names(x.shapes)
            elif x.has_text_frame and x.text_frame.text.strip() in ("Elsa Prada", "Pablo San-Jose", "César Robles"):
                for r in x.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(10)
    shrink_names(s.shapes)
    # drop the 3D sketch of the template, keep the funding logos of the layout
    pic = sh["Picture 24"]
    rid = pic._element.blipFill.find(qn("a:blip")).get(qn("r:embed"))
    pic._element.getparent().remove(pic._element)
    s.part.drop_rel(rid)
    textbox(s, 6.99, 0.50, 5.43, 0.75, ["<b>Thank you!</b>"], size=36)
    textbox(s, 6.99, 1.22, 5.43, 0.4, ["carlosp24.github.io · carlos.paya@csic.es"], size=16, color=RED)
    pub = lambda ref, what: (f"{ref} · <gry>{what}</gry>", {"bullet": True, "space": 1})
    textbox(s, 6.99, 1.70, 5.6, 3.4, [
        ("<b>PhD thesis papers</b> · full-shell nanowires", {"size": 13, "space": 3}),
        pub("PRB 107, 155423 (2023)", "CdGM analogs, theory"),
        pub("PRL 134, 206302 (2025)", "CdGM analogs, experiment"),
        pub("PRB 109, 115428 (2024)", "Majorana phenomenology"),
        pub("PRB 110, 115417 (2024)", "Absence of Majorana oscillations"),
        pub("PRB 111, 235420 (2025)", "Josephson effect"),
        pub("PRB 112, 134520 (2025)", "Fluxoid valve effect"),
        pub("PRL 137, 096606 (2026)", "this talk"),
        pub("arXiv:2607.18371 (2026)", "CdGMs in the critical current"),
        ("<b>Also during the PhD</b>", {"size": 13, "space": 3}),
        pub("SciPost Phys. 18, 069 (2025)", "Hole-based full shells"),
        pub("PRB 113, L161405 (2026)", "Non-Hermitian skin effect"),
    ], size=11, color=INK2)
    morph(s)
    notes(s, "Questions. Backup slides follow: models and parameters, electrostatics and skin width, origin of "
             "the skin for m_r = 0 vs m_r = 1, and the conductance with all m_J sectors and temperature.")


# ---------------------------------------------------------------------------
# Backups
# ---------------------------------------------------------------------------
def pdf_to_svg(pdf, out_dir):
    svg = out_dir / (pdf.stem + ".svg")
    png = out_dir / (pdf.stem + ".png")
    subprocess.run(["pdftocairo", "-svg", str(pdf), str(svg)], check=True)
    subprocess.run(["pdftocairo", "-png", "-r", "300", "-singlefile", str(pdf), str(out_dir / pdf.stem)],
                   check=True)
    return svg, png


def table(slide, x, y, w, rows, col_w, size=13, header=True):
    nr, nc = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(0.3 * nr))
    tb = gs.table
    for j, cw in enumerate(col_w):
        tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tb.cell(i, j)
            c.text = ""
            p = c.text_frame.paragraphs[0]
            fill_paragraph(p, val, size, INK, bold=(header and i == 0))
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string("F2F2F2" if (header and i == 0) else "FFFFFF")
    return gs


def backup_models(deck):
    s = deck.new("Backup: models and parameters")
    textbox(s, 0.92, 1.30, 11.5, 0.4, ["<b>Partial shell</b> (1D, Supplemental Material S-I)"], size=16)
    eq(s, "eq_h_partial", 1.1, 1.72, scale=0.8)
    eq(s, "eq_sigma_sc", 1.1, 2.35, scale=0.8)
    textbox(s, 0.92, 3.05, 11.5, 0.75, [
        "<b>Full shell</b> (Supplemental Material S-II): solid InAs-like core with a dome-shaped potential "
        "<i>U</i>(<i>r</i>), radial Rashba SOC, flux and Zeeman; the Al shell is integrated out into a "
        "Little–Parks (Abrikosov–Gor'kov) self-energy; solved per <i>m</i><sub>J</sub> sector"], size=14)
    table(s, 0.92, 3.95, 11.5, [
        ["", "Partial shell", "Full shell"],
        ["Common", "<i>m</i>* = 0.023 <i>m</i><sub>e</sub>, <i>Δ</i><sub>0</sub> = 0.23 meV",
         "<i>m</i>* = 0.023 <i>m</i><sub>e</sub>, <i>Δ</i><sub>0</sub> = 0.23 meV"],
        ["Filling, SOC", "<i>μ</i><sub>bulk</sub> = 2 meV, <i>α</i> = 40 meV nm",
         "<i>μ</i><sub>bulk</sub> = 22.8 meV, Δ<i>U</i> = 60 meV, ⟨<i>α</i>⟩ = 7 meV nm"],
        ["Coupling", "<i>Γ</i><sub>NS</sub> = 3<i>Δ</i><sub>0</sub>, <i>V</i><sub>C</sub> = 2<i>V</i><sub>Z</sub><sup>c</sup>",
         "<i>Γ</i><sub>NS</sub> = 40<i>Δ</i><sub>0</sub>, <i>g</i> = 10"],
        ["Geometry", "—", "<i>R</i> = 70 nm, <i>d</i> = 10 nm, <i>ξ</i><sub>d</sub> = 70 nm"],
        ["Smoothness (Fig. 1)", "<i>χ</i> = 200 nm", "<i>χ</i> = 1 µm"],
        ["Disorder (Fig. 3)", "<i>A</i><sub>imp</sub>/√<i>N</i> = 0.32<i>Δ</i><sub>0</sub>",
         "<i>A</i><sub>imp</sub>/√<i>N</i> = 0.02<i>Δ</i><sub>0</sub>"],
    ], [2.6, 4.2, 4.7], size=13)
    cite(s, [PAPER_SHORT + ", End Matter C and Supplemental Material"])
    notes(s, "Backup. Models and parameters (Appendix C of the paper).")


def backup_electrostatics(deck):
    s = deck.new("Backup: electrostatics and skin width")
    textbox(s, 0.92, 1.30, 6.0, 0.7, ["Laplace equation in the core, grounded shell, depleted end "
                                      "(<i>μ</i> = 0 at <i>z</i> = 0):"], size=15)
    eq(s, "eq_bessel", 1.0, 2.0, scale=0.9)
    textbox(s, 0.92, 3.0, 6.0, 0.7, ["Projected on the <i>m</i><sub>J</sub> = 0, <i>m</i><sub>r</sub> = 1 mode that "
                                     "turns topological:"], size=15)
    eq(s, "eq_mu_eff", 1.0, 3.75)
    table(s, 0.92, 4.45, 6.0, [
        ["Case", "Local phases along the wire"],
        ["Partial, <i>V</i><sub>Z</sub><sup>(2)</sup>",
         f"topological for <i>z</i> < {NUM['partial_topo_end_chi']:.2f}<i>χ</i>, trivial beyond"],
        ["Full, <i>Φ</i><sup>(1)</sup>", f"skin for <i>z</i> < {NUM['full_skin_end_chi_Phi1']:.2f}<i>χ</i>, then topological"],
        ["Full, <i>Φ</i><sup>(2)</sup>", f"skin for <i>z</i> < {NUM['full_skin_end_chi_Phi2']:.2f}<i>χ</i>, topological "
                                         f"up to {NUM['full_topo_end_chi_Phi2']:.2f}<i>χ</i>"],
    ], [1.9, 4.1], size=13)
    for k, n in enumerate(["s11_pd_full", "s11_pd_full_mzm", "s11_pd_full_qmzm", "s11_pd_full_skin"]):
        fig(s, n, 7.02, 1.22, descr=n.replace("_", " "))
    textbox(s, 7.12, 5.35, 5.3, 0.9, [f"Full-shell end signal below 1% of its maximum for <i>χ</i> > "
                                      f"{max(NUM['chi_vanish_full_MZM_nm'], NUM['chi_vanish_full_QMZM_nm']):.0f} nm "
                                      f"(MZM and quasi-MZM)"], size=14, color=INK2)
    cite(s, [PAPER_SHORT + ", End Matter A"])
    notes(s, "Backup. Electrostatics (End Matter A) and where the local phases change along the wire, from the "
             "phase diagrams and the projected potential.")


def backup_figure(deck, title, pdfs, note, ref, tmp):
    s = deck.new(title)
    # PNG only: the SVG conversion of these PDFs drops their transparency (grey bands turn black)
    from PIL import Image
    pngs = [pdf_to_svg(REPO / "plots" / "figures" / p, tmp)[1] for p in pdfs]
    wmax = 11.5 if len(pngs) == 1 else 5.6
    x = 0.92
    for png in pngs:
        W, H = (v / 300 for v in Image.open(png).size)
        sc = min(wmax / W, 5.4 / H)
        x0 = 0.92 + (11.5 - W * sc) / 2 if len(pngs) == 1 else x
        photo(s, png, x0, 1.32, w=W * sc, name=png.stem, descr=f"Paper figure {png.stem}")
        x += W * sc + 0.3
    cite(s, [ref])
    notes(s, note)


# ===========================================================================
def main(out):
    deck = Deck()
    prs = deck.prs
    set_theme_fonts(prs)
    slides = list(prs.slides)
    title_slide, farewell = slides[0], slides[29]
    for sl in slides[1:29]:
        drop_slide(prs, sl)

    slide_title(deck, title_slide)
    slide_who(deck)
    slide_majo_tunneling(deck)
    slide_smooth_partial(deck)
    slide_quasimajo(deck)
    slide_false_positive(deck)
    slide_fullshell(deck)
    slide_cdgm(deck)
    slide_cdgm_exp(deck)
    slide_mj(deck)
    slide_smooth_full(deck)
    slide_skin_origin(deck)
    slide_deep(deck)
    slide_no_fp(deck)
    slide_robust(deck)
    slide_conclusions(deck)
    slide_thanks(deck, farewell)
    main_slides = list(deck.order)
    deck.order = []
    with tempfile.TemporaryDirectory() as d:
        tmp = pathlib.Path(d)
        backup_models(deck)
        backup_electrostatics(deck)
        backup_figure(deck, "Backup: origin of the skin, <i>m</i><sub>r</sub> = 0 vs 1", ["Fig4.pdf"],
                      "Backup. Paper Fig. 4: band filling for the first and second radial modes, and the phase "
                      "diagram versus μ and ⟨α⟩ with the two thresholds. If asked about the axis: m_L here is the "
                      "eigenvalue of L_z = −i∂_φ + σ_z/2, half-integer (the m_J = 0 pair is m_L = ±1/2); the "
                      "'m_L ∈ ℤ' label of the published figure is a slip.", PAPER_SHORT + ", Fig. 4", tmp)
        backup_figure(deck, "Backup: conductance and temperature",
                      ["SFig_Conductance.pdf", "SFig_Temperature.pdf"],
                      "Backup. Supplemental Figs. S1 and S2: tunneling conductance including all m_J sectors "
                      "(T = 1 mK), and the m_J = 0 conductance at 10, 50 and 100 mK. Note for Q&A: the "
                      "conductance was computed at negative bias and mirrored, as in the Supplemental Material.",
                      PAPER_SHORT + ", Supplemental Material S-IV, S-V", tmp)
        backups = list(deck.order)

        # order: title, main slides, thank-you, backups
        lst = prs.slides._sldIdLst
        ids = {prs.part.related_part(sid.rId): sid for sid in lst}
        seq = [title_slide] + main_slides + [farewell] + backups
        for sid in list(lst):
            lst.remove(sid)
        for sl in seq:
            lst.append(ids[sl.part])
        # footer numbers follow the final order
        for k, sl in enumerate(prs.slides, start=1):
            for fld in sl._element.iter(qn("a:fld")):
                if fld.get("type") == "slidenum":
                    fld.find(qn("a:t")).text = str(k)
        prs.core_properties.title = "Absence of quasi-Majorana false positives in full-shell hybrid nanowires"
        prs.core_properties.author = "Carlos Payá"
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
    print(f"saved {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main(pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT)
